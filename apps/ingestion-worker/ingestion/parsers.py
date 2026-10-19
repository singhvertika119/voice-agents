"""Turn any supported input into the same thing: clean Markdown, plus any spreadsheet tables.

Format-specific mess stays in this file. Everything after it (chunking, embedding, search) never knows
what the original format was (docs/decision.md D-033).
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field


class ParseError(Exception):
    """Message is shown to the business owner as-is: say what happened and what to do."""


@dataclass
class Table:
    name: str
    columns: list[str]
    rows: list[dict]


@dataclass
class Parsed:
    markdown: str = ""
    tables: list[Table] = field(default_factory=list)
    kind: str = "text"  # text | faq


# ---------------------------------------------------------------- text / FAQ
_QA = re.compile(r"^\s*(?:Q(?:uestion)?\s*[:.)\-]|प्रश्न\s*[:.)\-])\s*(.+?)\s*$", re.I)
_ANS = re.compile(r"^\s*(?:A(?:nswer)?\s*[:.)\-]|उत्तर\s*[:.)\-])\s*(.*)$", re.I)


def parse_text(text: str) -> Parsed:
    text = text.replace("\r\n", "\n").strip()
    if not text:
        raise ParseError("The text is empty. Paste some content and try again.")
    return Parsed(markdown=text)


def parse_faq(text: str) -> Parsed:
    """Q: ... / A: ... pairs. Each pair becomes its own section so the question is the search key."""
    text = text.replace("\r\n", "\n").strip()
    pairs: list[tuple[str, list[str]]] = []
    for line in text.split("\n"):
        q, a = _QA.match(line), _ANS.match(line)
        if q:
            pairs.append((q.group(1), []))
        elif a and pairs:
            pairs[-1][1].append(a.group(1))
        elif pairs and line.strip():
            pairs[-1][1].append(line.strip())
    pairs = [(q, a) for q, a in pairs if " ".join(a).strip()]
    if not pairs:
        raise ParseError("No question and answer pairs found. Write each one as 'Q: ...' followed by 'A: ...'.")
    md = "\n\n".join(f"## {q}\n{' '.join(a).strip()}" for q, a in pairs)
    return Parsed(markdown=md, kind="faq")


# ---------------------------------------------------------------- PDF
def parse_pdf(data: bytes) -> Parsed:
    import fitz  # pymupdf

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ParseError("This PDF could not be opened. Upload a different copy.") from e
    pages, empty_pages = [], []
    for n, page in enumerate(doc, start=1):
        # Repeated headers/footers: drop lines in the top/bottom 6% of the page.
        h = page.rect.height
        blocks = [b for b in page.get_text("blocks") if b[6] == 0 and h * 0.06 < b[1] < h * 0.94]
        blocks.sort(key=lambda b: (round(b[1] / 8), b[0]))
        text = "\n".join(b[4].strip() for b in blocks).strip()
        if len(text) < 20:
            empty_pages.append(n)
        else:
            pages.append(text)
    if not pages:
        raise ParseError("This PDF has no readable text. It looks like a scanned image. Scanned PDFs are not supported yet; upload a text PDF or paste the text.")
    md = "\n\n".join(pages)
    if _looks_like_legacy_hindi(md):
        raise ParseError("This PDF uses an old Hindi font that comes out as garbled text. Paste the text instead, or re-save it as a new PDF.")
    note = f"\n\n(Pages without text were skipped: {', '.join(map(str, empty_pages))}.)" if empty_pages else ""
    return Parsed(markdown=_promote_headings(md) + note)


def _looks_like_legacy_hindi(text: str) -> bool:
    """Kruti Dev and similar fonts extract as Latin gibberish with no Devanagari. Heuristic: lots of
    punctuation-heavy 'words' and no Devanagari, with accented/symbol characters typical of those fonts."""
    sample = text[:4000]
    if re.search(r"[ऀ-ॿ]", sample):
        return False
    letters = sum(c.isalpha() for c in sample) or 1
    odd = sum(c in "ìíîïðñòóôõöøùúûüýþÿ¡¢£¤¥¦§¨©ª«¬®¯°±²³´µ¶·¸¹º»¼½¾¿" for c in sample)
    return odd / letters > 0.05


def _promote_headings(text: str) -> str:
    """PDFs have no heading markup. Treat a short line without final punctuation, followed by a longer
    line, as a heading so chunking can follow document structure."""
    out, lines = [], text.split("\n")
    for i, line in enumerate(lines):
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        s = line.strip()
        if s and len(s) <= 60 and not s.endswith((".", ",", ";", ":", "?", "।")) and len(nxt.strip()) > len(s) and (s.isupper() or s.istitle()):
            out.append(f"## {s}")
        else:
            out.append(line)
    return "\n".join(out)


# ---------------------------------------------------------------- DOCX
def parse_docx(data: bytes) -> Parsed:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError("This Word file could not be opened. Save it as .docx and upload again.") from e
    out: list[str] = []
    body = d.element.body
    paras = {p._p: p for p in d.paragraphs}
    tables = {t._tbl: t for t in d.tables}
    for el in body.iterchildren():
        if el in paras:
            p = paras[el]
            txt = p.text.strip()
            if not txt:
                continue
            m = re.match(r"Heading\s*(\d)", p.style.name or "")
            out.append(f"{'#' * min(int(m.group(1)), 4)} {txt}" if m else txt)
        elif el in tables:
            out.append(_table_to_sentences([[c.text.strip() for c in r.cells] for r in tables[el].rows]))
    md = "\n\n".join(x for x in out if x)
    if not md.strip():
        raise ParseError("This Word file has no text in it.")
    return Parsed(markdown=md)


def _table_to_sentences(rows: list[list[str]]) -> str:
    """Voice cannot read a table, so each row becomes a sentence-like line: 'Service: ECG. Price: 400.'"""
    rows = [r for r in rows if any(c for c in r)]
    if len(rows) < 2:
        return " ".join(" ".join(r) for r in rows)
    head = rows[0]
    return "\n".join(". ".join(f"{h}: {c}" for h, c in zip(head, r) if c) + "." for r in rows[1:])


# ---------------------------------------------------------------- spreadsheets
def parse_csv(data: bytes, name: str = "Sheet") -> Parsed:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    return _table_parsed(name, rows)


def parse_xlsx(data: bytes) -> Parsed:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ParseError("This spreadsheet could not be opened. Save it as .xlsx or .csv and upload again.") from e
    result = Parsed(kind="text")
    for ws in wb.worksheets:
        rows = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if len(rows) >= 2:
            sub = _table_parsed(ws.title, rows)
            result.tables += sub.tables
    if not result.tables:
        raise ParseError("No table with a header row and data rows was found in this spreadsheet.")
    result.markdown = _tables_overview(result.tables)
    return result


def _table_parsed(name: str, rows: list[list[str]]) -> Parsed:
    if len(rows) < 2:
        raise ParseError("The spreadsheet needs a header row and at least one row of data.")
    cols = [c.strip() or f"Column {i + 1}" for i, c in enumerate(rows[0])]
    data = [{c: (r[i].strip() if i < len(r) else "") for i, c in enumerate(cols)} for r in rows[1:]]
    t = Table(name=name, columns=cols, rows=data)
    return Parsed(markdown=_tables_overview([t]), tables=[t])


def _tables_overview(tables: list[Table]) -> str:
    """A short description so the agent knows which lookup tables exist; the rows themselves are not embedded."""
    return "\n\n".join(
        f"## Table: {t.name}\nColumns: {', '.join(t.columns)}. {len(t.rows)} rows. Use the lookup tool for exact values."
        for t in tables
    )


# ---------------------------------------------------------------- web page
def parse_html(html: str) -> Parsed:
    import trafilatura

    # favor_recall keeps FAQ and secondary sections that favor_precision drops (measured on a real site: +35% text).
    md = trafilatura.extract(html, output_format="markdown", include_tables=True, include_links=False,
                             include_formatting=True, favor_recall=True)
    if not md or len(md.strip()) < 80:
        raise ParseError("Not enough readable text was found on this page. It may need JavaScript to load, or may be blocked. Paste the text instead.")
    return Parsed(markdown=_restore_headings(md, html))


def _restore_headings(md: str, html: str) -> str:
    """trafilatura often flattens headings to plain lines. Chunking relies on headings for section names,
    so mark any line that matches an <h1>-<h3> from the original page."""
    import lxml.html

    try:
        tree = lxml.html.fromstring(html)
    except Exception:
        return md
    norm = lambda t: re.sub(r"\s+", " ", t).strip().lower()  # noqa: E731
    heads = {}
    for level, tag in enumerate(("h1", "h2", "h3"), start=1):
        for el in tree.iter(tag):
            for br in el.iter("br"):          # "Line one<br>line two" is one heading: join with a space
                br.tail = " " + (br.tail or "")
            text = norm(el.text_content())
            if 3 <= len(text) <= 120:
                heads.setdefault(text, level)
    out = []
    for line in md.split("\n"):
        plain = norm(re.sub(r"^#+\s*|[*_]", "", line))
        level = heads.get(plain)
        if level and line.strip() and not line.startswith("#"):
            text = re.sub(r"[*_]", "", line).strip()
            line = "#" * level + " " + text
        out.append(line)
    return "\n".join(out)


def parse_file(kind: str, data: bytes, name: str) -> Parsed:
    kind = kind.lower()
    if kind == "pdf":
        return parse_pdf(data)
    if kind == "docx":
        return parse_docx(data)
    if kind == "csv":
        return parse_csv(data, name.rsplit(".", 1)[0])
    if kind == "xlsx":
        return parse_xlsx(data)
    if kind in ("txt", "md"):
        return parse_text(data.decode("utf-8-sig", errors="replace"))
    raise ParseError(f"Files of type .{kind} are not supported. Use PDF, Word, Excel, CSV or text.")

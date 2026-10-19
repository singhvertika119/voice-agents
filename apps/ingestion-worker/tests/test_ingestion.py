import io

import pytest

from ingestion import parsers
from ingestion.chunking import chunk_markdown, normalize_for_voice
from ingestion.netsafe import assert_public
from ingestion.parsers import ParseError


def test_faq_pairs_become_sections():
    p = parsers.parse_faq("Q: Clinic kab khulta hai?\nA: Subah 9 se shaam 7 baje.\n\nQ: Fees?\nA: Rs 600.")
    assert p.kind == "faq"
    chunks = chunk_markdown(p.markdown)
    assert [c.section for c in chunks] == ["Clinic kab khulta hai?", "Fees?"]
    assert "600 rupees" in chunks[1].content


def test_faq_without_pairs_gives_readable_error():
    with pytest.raises(ParseError, match="Q: "):
        parsers.parse_faq("just some text")


def test_chunks_follow_headings_and_carry_section_path():
    md = "# Refunds\n\n## Timelines\nRefunds take 5 to 7 days.\n\n## Eligibility\nOnly unused items qualify."
    chunks = chunk_markdown(md)
    assert chunks[0].section == "Refunds > Timelines"
    assert chunks[0].content.startswith("Refunds > Timelines: ")
    assert chunks[1].section == "Refunds > Eligibility"


def test_long_text_is_split_into_bounded_chunks():
    md = "## Policy\n" + " ".join(f"Sentence number {i} is here." for i in range(200))
    chunks = chunk_markdown(md)
    assert len(chunks) > 1 and all(len(c.content) < 1700 for c in chunks)


def test_voice_normalisation_drops_urls_and_citations():
    out = normalize_for_voice("See https://x.com/a for more [1]. Fee is ₹1,250 **only**.")
    assert "http" not in out and "[1]" not in out and "1250 rupees" in out and "*" not in out


def test_csv_becomes_lookup_table_not_prose():
    p = parsers.parse_csv("Service,Price\nECG,400\nX-ray,700\n".encode(), "Prices")
    t = p.tables[0]
    assert t.columns == ["Service", "Price"] and t.rows[1] == {"Service": "X-ray", "Price": "700"}
    assert "400" not in p.markdown  # values are not embedded, only described


def test_csv_needs_data_rows():
    with pytest.raises(ParseError):
        parsers.parse_csv(b"Service,Price\n")


def test_docx_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Fees", 1)
    d.add_paragraph("Our fee schedule.")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Service", "Price"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "ECG", "400"
    buf = io.BytesIO()
    d.save(buf)
    md = parsers.parse_docx(buf.getvalue()).markdown
    assert md.startswith("# Fees") and "Service: ECG. Price: 400." in md


def test_xlsx_sheets_become_tables():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fees"
    ws.append(["Class", "Fee"])
    ws.append(["Nursery", 12000])
    buf = io.BytesIO()
    wb.save(buf)
    p = parsers.parse_xlsx(buf.getvalue())
    assert p.tables[0].name == "Fees" and p.tables[0].rows[0]["Fee"] == "12000"


def _pdf(text: str | None) -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_textbox(fitz.Rect(72, 100, 520, 700), text, fontsize=11)
    return doc.tobytes()


def test_digital_pdf_text_extracted():
    p = parsers.parse_pdf(_pdf("Admissions\nAdmissions open in April every year for classes one to eight at our school."))
    assert "Admissions open in April" in p.markdown


def test_scanned_pdf_gives_readable_error():
    with pytest.raises(ParseError, match="scanned"):
        parsers.parse_pdf(_pdf(None))


def test_html_main_content_only():
    html = "<html><body><nav>Home About</nav><article><h1>Our clinic</h1>" + "<p>We are open Monday to Saturday from nine to seven and offer cardiology and pediatric care.</p>" * 3 + "</article></body></html>"
    assert "open Monday to Saturday" in parsers.parse_html(html).markdown


@pytest.mark.parametrize("url", ["http://127.0.0.1/admin", "http://localhost:8000", "http://169.254.169.254/latest/meta-data", "ftp://example.com", "http://10.0.0.5"])
def test_private_addresses_are_blocked(url):
    with pytest.raises(ParseError):
        assert_public(url)


def test_unsupported_file_type():
    with pytest.raises(ParseError, match="not supported"):
        parsers.parse_file("exe", b"", "a.exe")


def test_html_headings_are_restored_for_section_names():
    body = "<p>We are open Monday to Saturday from nine to seven and offer cardiology and pediatric care.</p>" * 2
    html = f"<html><body><article><h1>Our clinic</h1>{body}<h2>Fees</h2><p>Consultation costs six hundred rupees and a follow up visit costs three hundred rupees.</p></article></body></html>"
    chunks = chunk_markdown(parsers.parse_html(html).markdown)
    assert any(c.section.endswith("Fees") for c in chunks)
